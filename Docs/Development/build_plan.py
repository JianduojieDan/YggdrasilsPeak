from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

ROOT = Path(__file__).parent
source = ROOT / '脊梁世界树之巅开发规划.md'
doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21)
sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
    st = doc.styles[name]
    st.font.name = 'Microsoft YaHei'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    st.font.color.rgb = RGBColor(0,0,0)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.1
for name,size in [('Title',23),('Heading 1',17),('Heading 2',12.5),('Heading 3',11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
for st in doc.styles:
    if st.element.pPr is not None:
        for border in list(st.element.pPr.findall(qn('w:pBdr'))):
            st.element.pPr.remove(border)
header = sec.header.paragraphs[0]
header.text = '脊梁 世界树之巅    开发规划与任务清单'
header.runs[0].font.size = Pt(8)
footer = sec.footer.paragraphs[0]; footer.alignment = 2
footer.add_run('2026年10月5日   |   ')
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
active_section = ''
for line in source.read_text(encoding='utf-8').splitlines():
    if not line.strip(): continue
    if line == '---':
        doc.add_page_break(); continue
    if line.startswith('# '): doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        active_section = line[3:5]
        doc.add_paragraph(line[3:], 'Heading 1')
    elif line.startswith('### '): doc.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); p.add_run(line[2:]); p.paragraph_format.keep_together = True
    else: doc.add_paragraph(line)
    if active_section == '05':
        doc.paragraphs[-1].paragraph_format.line_spacing = 1.0
        doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
doc.core_properties.title = '脊梁世界树之巅开发规划与任务清单'
doc.core_properties.subject = '基于游戏设定和当前 Unreal Engine 项目的开发路线'
doc.core_properties.author = '项目开发规划'
out=ROOT/'脊梁世界树之巅开发规划.docx'
doc.save(out)
print(json.dumps({'output':str(out),'paragraphs':len(doc.paragraphs),'characters':sum(len(p.text) for p in doc.paragraphs)},ensure_ascii=False))
